import base64
from io import BytesIO
import re

from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.contract_apuracao import seed_contract_defaults
from app.database import Base, get_db
from app.glpi_contract import GlpiSnapshot, evaluate
from app.models import ContractAbsence, ContractInvoiceMap, ContractRegistryStatus, User
from app.routes import contract, contract_registry
from app.security import create_session_token, generate_csrf_token, settings
from app.timezone import format_local_datetime
from tests.test_contract_rules import SLAS, ticket


@pytest.fixture
def client_factory(monkeypatch):
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as db:
        admin = User(username='contract_admin', is_admin=True, is_active=True)
        db.add(admin)
        db.commit()
        token = create_session_token(admin.id)
        seed_contract_defaults(db)
    application = FastAPI()
    application.state.settings = settings
    application.state.templates = Jinja2Templates(directory='templates')
    application.state.templates.env.filters['local_datetime'] = format_local_datetime
    application.state.templates.env.globals['menu_dashboards_for_user'] = lambda _: []
    application.mount('/static', StaticFiles(directory='static'), name='static')
    application.include_router(contract.router)
    application.include_router(contract_registry.router)

    def session():
        with factory() as db:
            yield db
    application.dependency_overrides[get_db] = session

    @application.middleware('http')
    async def csrf_context(request, call_next):
        request.state.csrf_token = generate_csrf_token(request.cookies.get(settings.cookie_name, ''))
        return await call_next(request)

    def fake_glpi(db, period, cut, holidays, entity_invoice):
        snapshot = GlpiSnapshot(tickets=[
            ticket(1, '2026-09-22 08:00:00', '2026-09-22 15:00:00', ttr='2026-09-22 12:00:00'),
            ticket(2, '2026-10-01 08:00:00', '2026-10-01 10:00:00', ttr='2026-10-01 12:00:00', entity=7),
            ticket(3, '2026-10-05 08:00:00', '2026-10-07 10:00:00', entity=7, priority=3),
        ], logs=[], slas=SLAS, cut=cut)
        return evaluate(snapshot, period, holidays, entity_invoice), ['glpi_knowbaseitems_items'], None
    monkeypatch.setattr('app.contract_apuracao.load_glpi', fake_glpi)
    with TestClient(application, follow_redirects=False) as client:
        client.cookies.set(settings.cookie_name, token)
        yield client, token, factory
    engine.dispose()


def post(client, token, url, **data):
    return client.post(url, data={'csrf_token': generate_csrf_token(token), 'competencia': '2026-10', **data})


def test_registry_flow_import_confirm_and_reopen(client_factory):
    client, token, factory = client_factory
    page = client.get('/admin/contract/cadastros?competencia=2026-10')
    assert page.status_code == 200, page.text
    assert 'Indicador → fatura' in page.text and 'operador_teleatendimento' in page.text
    raw = 'data;equipe;perfil;pessoa;coberta;observacao\n01/10/2026;1;operador_teleatendimento;M1;não;\n'.encode()
    response = post(client, token, '/admin/contract/cadastros/importar', cadastro='ausencias', csv_base64=base64.b64encode(raw).decode())
    assert response.status_code == 303
    assert post(client, token, '/admin/contract/cadastros/confirmar', cadastro='ausencias_item01').status_code == 303
    with factory() as db:
        assert db.scalar(select(ContractAbsence)).person_ref == 'M1'
        assert db.scalar(select(ContractRegistryStatus)).registry == 'ausencias_item01'
    page = client.get('/admin/contract?competencia=2026-10&corte=2026-10-10T12:00')
    ins11 = re.search(r'data-indicator="INS11" data-scope="item01" data-priority="" data-value="([^"]*)"', page.text)
    assert ins11 and ins11.group(1).startswith('99.6')
    assert post(client, token, '/admin/contract/cadastros/importar', cadastro='ausencias',
                csv_content=raw.decode()).status_code == 303
    with factory() as db:
        assert db.scalar(select(ContractRegistryStatus)) is None
    bad = post(client, token, '/admin/contract/cadastros/importar', cadastro='ausencias',
               csv_content='data;equipe;perfil;pessoa;coberta;observacao\n01/11/2026;1;operador_teleatendimento;M1;não;\n')
    assert bad.status_code == 400 and 'fora da competência' in bad.text


def test_invoice_confirmation_requires_source(client_factory):
    client, token, factory = client_factory
    refused = post(client, token, '/admin/contract/cadastros/fatura', indicador='INS02', fatura='item01', confirmado='1', fonte='')
    assert refused.status_code == 400
    assert post(client, token, '/admin/contract/cadastros/fatura', indicador='INS02', fatura='item01_geral', confirmado='1',
                fonte='Resposta da contratada em 18/10/2026').status_code == 303
    with factory() as db:
        row = db.get(ContractInvoiceMap, 'INS02')
        assert (row.invoice, row.confirmed) == ('item01_geral', True)


def test_holiday_and_entity_forms(client_factory):
    client, token, factory = client_factory
    assert post(client, token, '/admin/contract/cadastros/feriado', dia='2027-10-24', nome='Aniversário de Goiânia',
                tipo='municipal', fonte='Lei municipal').status_code == 303
    assert post(client, token, '/admin/contract/cadastros/feriado', dia='2027-10-24', tipo='xx', nome='a', fonte='b').status_code == 400
    assert post(client, token, '/admin/contract/cadastros/entidade', entidade='9', fatura='item02', fonte='teste').status_code == 303
    assert post(client, token, '/admin/contract/cadastros/entidade', entidade='9', acao='remover').status_code == 303
    assert post(client, token, '/admin/contract/cadastros/entidade', entidade='9', acao='remover').status_code == 400


def test_full_apuracao_page_and_stable_export(client_factory):
    client, token, factory = client_factory
    page = client.get('/admin/contract?competencia=2026-10&corte=2026-10-21T18:00')
    assert page.status_code == 200, page.text
    assert 'Apuração completa da competência 2026-10' in page.text
    assert 'data-invoice="item01"' in page.text and 'data-invoice="item02"' in page.text
    assert 'TRS45: 1 chamados' in page.text
    assert 'glpi_knowbaseitems_items' in page.text
    url = re.search(r'href="(/admin/contract/export[^\"]+formato=xlsx)"', page.text).group(1).replace('&amp;', '&')
    assert 'corte=2026-10-21T18%3A00' in url
    again = client.get('/admin/contract?competencia=2026-10&corte=2026-10-21T18:00')
    assert re.search(r'assinatura=([0-9a-f]{64})', again.text).group(1) == re.search(r'assinatura=([0-9a-f]{64})', url).group(1)
    response = client.get(url)
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ['Memória da apuração', 'Indicadores', 'Fora do prazo', 'TRS45', 'Reaberturas', 'Diagnósticos', 'Pendências']
    assert [c.value for c in book['TRS45'][2]][:2] == ['3', 'item02']
    assert ('INS03', 'item01', 'P3', '1') in {tuple(c.value for c in r) for r in book['Fora do prazo'].iter_rows(min_row=2)}
    json_url = url.replace('formato=xlsx', 'formato=json')
    evidence = client.get(json_url).json()
    assert evidence['apuracao']['cut'] == '2026-10-21T18:00:00'
    assert any(p.startswith('P1+P2: Excedentes') for p in evidence['pending'])
    assert not any(p.startswith('TRS45:') or p.startswith('INS01:') for p in evidence['pending'])
    moved = client.get(json_url.replace('corte=2026-10-21T18%3A00', 'corte=2026-10-21T19%3A00'))
    assert moved.status_code == 409


def test_ura_replacement_through_form(client_factory):
    from tests.test_ura_import import sample_csv
    client, token, factory = client_factory
    def upload(raw, **extra):
        return client.post('/admin/contract/ura', data={'csrf_token': generate_csrf_token(token), 'competencia': '2026-10',
                                                        'projeto': 'DETRAN-GO', 'csv_base64': base64.b64encode(raw).decode(), **extra})
    assert upload(sample_csv()).status_code == 303
    assert upload(sample_csv(answered=9)).status_code == 400
    done = upload(sample_csv(answered=9), substituir='1', justificativa='Reenvio corrigido pela contratada')
    assert done.status_code == 303
    page = client.get(done.headers['location'])
    assert '1 substituídos (justificativa: Reenvio corrigido pela contratada)' in page.text


def test_staff_reimport_reopens_every_confirmed_competence(client_factory):
    client, token, factory = client_factory
    raw = 'pessoa;equipe;perfil;inicio;fim;desligamento\nA;1;operador_teleatendimento;22/09/2026;;\n'
    assert post(client, token, '/admin/contract/cadastros/importar', cadastro='pessoal', csv_content=raw).status_code == 303
    confirm = client.post('/admin/contract/cadastros/confirmar', data={'csrf_token': generate_csrf_token(token), 'competencia': '2026-12', 'cadastro': 'pessoal'})
    assert confirm.status_code == 303
    assert post(client, token, '/admin/contract/cadastros/importar', cadastro='pessoal', csv_content=raw).status_code == 303
    with factory() as db:
        assert db.scalars(select(ContractRegistryStatus)).all() == []

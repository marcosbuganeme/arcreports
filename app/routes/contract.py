"""Apuração contratual administrativa, separada dos relatórios SQL instalados."""
import base64
import binascii
import csv
import hashlib
import json
from io import BytesIO, StringIO
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Request
from fastapi.responses import RedirectResponse, JSONResponse, Response
from openpyxl import Workbook
from sqlalchemy.orm import Session
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy import select

from app.billing_rules import Competence, contract_month, pao_factor, decimal_value, adjusted_goal, weighted_points, knowledge_goal, staffing_window, previous_competence, current_competence
from app.database import get_db
from app import contract_apuracao
from app.contract_report import ura_report, calculation_evidence, evidence_rows
from app.models import UraImportRun
from app.ura import import_ura_csv, normalize_project, MAX_CSV_BYTES
from app.routes.common import render, form_data
from app.security import require_admin, verify_csrf

router = APIRouter()


def contract_context(params):
    period = Competence.parse(params['competencia']) if params.get('competencia') else current_competence()
    base_goal = decimal_value(params.get('meta_base') or '95')
    direction = params.get('direcao') or 'minimum'
    base_points = decimal_value(params.get('pontos_base') or '0')
    return {'period': period, 'previous_period':previous_competence(period), 'contract_month': contract_month(period), 'factor': pao_factor(period),
            'staffing_window': staffing_window(period), 'knowledge_goal': knowledge_goal(period), 'trimester': (contract_month(period) - 1) // 3 + 1,
            'base_points': base_points, 'weighted_points': weighted_points(base_points, period),
            'base_goal': base_goal, 'direction': direction, 'goal': adjusted_goal(base_goal, direction, period)}


def report_context(db, params):
    context = contract_context(params)
    context['project'] = normalize_project(params.get('projeto') or 'DETRAN-GO')
    service = params.get('servico') or ''
    try:
        service_id = int(service) if service else None
    except ValueError as exc:
        raise ValueError('Código de serviço inválido.') from exc
    if service_id is not None and (service_id < 0 or service_id > 2147483647):
        raise ValueError('Código de serviço inválido.')
    context['ura'] = ura_report(db, context['period'], context['project'], service_id)
    contract_month(context['period'])
    cut = contract_apuracao.parse_cut(params.get('corte'), context['period'], contract_apuracao.portal_now())
    context['apuracao'] = contract_apuracao.build(db, context['period'], context['project'], cut)
    context['evidence'] = calculation_evidence(context)
    context['evidence']['staffing']['result'] = [r['state'] for r in context['apuracao']['rows'] if r['code'] in ('INS09', 'INS10')]
    context['evidence']['staffing']['rule'] = 'INS09/10: desde o mês 3; três competências; desligados sem justa causa ÷ alocados distintos'
    context['evidence']['pending'] = [f'{a}: {b}' for a, b in context['apuracao']['pending']]
    context['evidence']['apuracao'] = contract_apuracao.evidence(context['apuracao'])
    context['invoice_labels'] = contract_apuracao.INVOICES
    context['invoice_choices'] = contract_apuracao.INVOICE_CHOICES
    context['scope_labels'] = contract_apuracao.SCOPE_LABELS
    canonical = json.dumps(context['evidence'], ensure_ascii=False, sort_keys=True, separators=(',',':'))
    context['digest'] = hashlib.sha256(canonical.encode('utf-8')).hexdigest()
    context['export_query'] = urlencode({'competencia':context['period'].label,'projeto':context['project'],
        'servico':service or '', 'meta_base':str(context['base_goal']), 'direcao':context['direction'],
        'pontos_base':str(context['base_points']), 'corte':cut.isoformat(timespec='minutes'), 'assinatura':context['digest']})
    return context


@router.get('/admin/contract')
async def contract_page(request: Request, db: Session = Depends(get_db)):
    user = require_admin(request, db)
    if isinstance(user, RedirectResponse):
        return user
    try:
        context = report_context(db, request.query_params)
    except ValueError as exc:
        return render(request, 'contract.html', {'active': 'contract', 'error': str(exc)}, db, status_code=400)
    context['imports'] = db.scalars(select(UraImportRun).where(UraImportRun.project == context['project']).order_by(UraImportRun.id.desc()).limit(20)).all()
    context['imported_run'] = next((run for run in context['imports'] if str(run.id) == request.query_params.get('importacao')), None)
    return render(request, 'contract.html', {'active': 'contract', **context}, db)


@router.post('/admin/contract/ura', dependencies=[Depends(verify_csrf)])
async def upload_ura(request: Request, db: Session = Depends(get_db)):
    user = require_admin(request, db)
    if isinstance(user, RedirectResponse):
        return user
    data = await form_data(request)
    try:
        period = Competence.parse(data['competencia']) if data.get('competencia') else current_competence()
        contract_month(period)
        project = normalize_project(data.get('projeto'))
        encoded = data.get('csv_base64') or ''
        if len(encoded) > (MAX_CSV_BYTES * 4 // 3 + 8):
            raise ValueError('O CSV deve ter no máximo 2 MB.')
        try:
            raw = base64.b64decode(encoded, validate=True) if encoded else data.get('csv_content','').encode('utf-8')
        except binascii.Error as exc:
            raise ValueError('Arquivo inválido. Selecione o CSV novamente.') from exc
        run = import_ura_csv(db, raw, project=project, source_name=data.get('source_name') or 'ura.csv', username=user.username,
                             replace=data.get('substituir') == '1', reason=data.get('justificativa'))
    except ValueError as exc:
        return render(request, 'contract.html', {'active':'contract', 'error':str(exc)}, db, status_code=400)
    except SQLAlchemyError:
        db.rollback()
        return render(request, 'contract.html', {'active':'contract', 'error':'Falha ao gravar a importação. Nenhum registro deste arquivo foi salvo. Confira o banco local e tente novamente.'}, db, status_code=409)
    query = urlencode({'competencia':period.label, 'projeto':project, 'importacao':run.id})
    return RedirectResponse('/admin/contract?' + query, status_code=303)


def detail_sheets(apuracao):
    """Abas de conferência: cada chamado que pesou na apuração, para auditoria linha a linha."""
    indicators = [('indicador', 'escopo', 'prioridade', 'situação', 'resultado', 'numerador', 'denominador', 'meta', 'desvio',
                   'pontos proporcional', 'pontos truncado', 'observação')]
    for r in apuracao['rows']:
        points = r['points'] or {}
        indicators.append((r['code'], r['scope'], r['priority'], r['state'], r['value'], r['numerator'], r['denominator'], r['goal'],
                           points.get('deviation'), points.get('proporcional'), points.get('truncado'),
                           '; '.join(filter(None, [r['note'], points.get('pending'), *r['missing']]))))
    late = [('indicador', 'escopo', 'prioridade', 'chamado')]
    glpi = apuracao['glpi']
    if glpi:
        for scope, data in glpi['scopes'].items():
            for code in ('INS02', 'INS03', 'INS04'):
                for priority, item in data[code].items():
                    late += [(code, scope, priority, ticket) for ticket in item['late']]
    trs = [('chamado', 'fatura', 'prioridade', 'tipo', 'zabbix', 'abertura', 'solução', 'SLA', 'limite (s)', 'tempo líquido (s)', 'situação')]
    trs += [(t['ticket'], t['scope'], t['priority'], t['type'], 'sim' if t['zabbix'] else 'não', t['opened'], t['solved'], t['sla'],
             t['limit_seconds'], t['net_seconds'], t['state']) for t in apuracao['trs45']]
    reopen = [('chamado', 'prioridade', 'conclusão', 'reabertura')]
    reopen += [(e['ticket'], e['priority'], e['conclusion'], e['reopened']) for e in (glpi['reopen_events'] if glpi else [])]
    diagnostics = [('diagnóstico', 'chamados')]
    if glpi:
        for key, value in glpi['diagnostics'].items():
            diagnostics.append((key, ', '.join(map(str, value)) if isinstance(value, list) else value))
    pending = [('assunto', 'pendência')] + list(apuracao['pending'])
    return [('Indicadores', indicators), ('Fora do prazo', late), ('TRS45', trs), ('Reaberturas', reopen),
            ('Diagnósticos', diagnostics), ('Pendências', pending)]


@router.get('/admin/contract/export')
async def export_contract(request: Request, db: Session = Depends(get_db)):
    user = require_admin(request, db)
    if isinstance(user, RedirectResponse):
        return user
    try:
        context = report_context(db, request.query_params)
    except ValueError as exc:
        return JSONResponse({'error':str(exc)}, status_code=400)
    if request.query_params.get('assinatura') != context['digest']:
        return JSONResponse({'error':'A apuração mudou ou a assinatura está ausente. Atualize a tela e exporte novamente.'},status_code=409)
    evidence = context['evidence']
    kind = request.query_params.get('formato','json')
    filename = f"apuracao-{context['project']}-{context['period'].label}"
    headers = {'Content-Disposition': f'attachment; filename="{filename}.{kind if kind in {"json","csv","xlsx"} else "txt"}"',
               'X-Apuracao-SHA256':context['digest']}
    if kind == 'json':
        return JSONResponse(evidence,headers=headers)
    rows = [('campo','valor'),('apuracao_sha256',context['digest']),*evidence_rows(evidence)]
    if kind == 'csv':
        output = StringIO(newline='')
        writer = csv.writer(output,delimiter=';')
        for row in rows:
            writer.writerow(["'"+cell if cell.startswith(('=','+','-','@','\t','\r','\n')) else cell for cell in row])
        return Response(('\ufeff'+output.getvalue()).encode('utf-8'),media_type='text/csv; charset=utf-8',headers=headers)
    if kind == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = 'Memória da apuração'
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                cell.data_type = 's'  # Conteúdo importado nunca vira fórmula de planilha.
        sheet.column_dimensions['A'].width = 45
        sheet.column_dimensions['B'].width = 85
        sheet.freeze_panes = 'A2'
        for title, table in detail_sheets(context['apuracao']):
            extra = book.create_sheet(title)
            for row in table:
                extra.append(['' if v is None else str(v) for v in row])
                for cell in extra[extra.max_row]:
                    cell.data_type = 's'
            extra.freeze_panes = 'A2'
        output = BytesIO()
        book.save(output)
        return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers=headers)
    return JSONResponse({'error':'Formato inválido. Use json, csv ou xlsx.'},status_code=400)
